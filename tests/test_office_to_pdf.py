from io import BytesIO

from openpyxl import Workbook
from pptx import Presentation
import pytest

import app.main as main
from conftest import page_layout


needs_libreoffice = pytest.mark.skipif(main.find_soffice() is None, reason="LibreOffice is not installed")


def make_xlsx(sheet_count: int) -> bytes:
    workbook = Workbook()
    workbook.active.title = "Hoja 1"
    for number in range(2, sheet_count + 1):
        workbook.create_sheet(f"Hoja {number}")
    for number, sheet in enumerate(workbook.worksheets, start=1):
        sheet["A1"] = f"Hoja {number}"
        sheet["B1"] = number * 10
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_pptx(slide_count: int) -> bytes:
    presentation = Presentation()
    for number in range(1, slide_count + 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"Diapositiva {number}"
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    ("endpoint", "filename"),
    [
        ("/api/excel-to-pdf", "notes.docx"),
        ("/api/excel-to-pdf", "data.csv"),
        ("/api/powerpoint-to-pdf", "book.xlsx"),
        ("/api/powerpoint-to-pdf", "slides.key"),
    ],
)
def test_office_rejects_other_extensions(client, endpoint, filename):
    response = client.post(endpoint, files={"file": (filename, b"hello")})

    assert response.status_code == 400


@pytest.mark.parametrize(
    ("endpoint", "filename"),
    [("/api/excel-to-pdf", "a.xlsx"), ("/api/powerpoint-to-pdf", "a.pptx")],
)
def test_office_reports_missing_libreoffice(client, monkeypatch, endpoint, filename):
    monkeypatch.setattr(main, "find_soffice", lambda: None)

    response = client.post(endpoint, files={"file": (filename, b"content")})

    assert response.status_code == 503


@needs_libreoffice
def test_excel_converts_every_sheet(client):
    response = client.post(
        "/api/excel-to-pdf",
        files={"file": ("book.xlsx", make_xlsx(2))},
        data={"output_name": "informe"},
    )

    assert response.status_code == 200
    assert response.headers["content-disposition"] == 'attachment; filename="informe.pdf"'
    assert len(page_layout(response.content)) == 2


@needs_libreoffice
def test_powerpoint_converts_one_page_per_slide(client):
    response = client.post("/api/powerpoint-to-pdf", files={"file": ("deck.pptx", make_pptx(3))})

    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    assert len(page_layout(response.content)) == 3

